import openpyxl
import pyexcelerate as px
import polars as pl

def open_template(template_path):
    return openpyxl.load_workbook(template_path)


def dump_data_to_sheet(
    data=None,
    sheet_name="data",
    excel_thing=None,
    starting_cell="A1",
    write_header=True,
):
    if sheet_name not in excel_thing.sheetnames:
        excel_thing.create_sheet(sheet_name)

    sheet = excel_thing[sheet_name]
    start_row = sheet[starting_cell].row
    start_col = sheet[starting_cell].column

    if write_header:
        for index, col in enumerate(data.columns, start=start_col):
            sheet.cell(row=start_row, column=index).value = col
    index = 0
    for row in data.iter_rows(named=True):
        for col_index, col in enumerate(data.columns, start=start_col):
            sheet.cell(
                row=start_row + index,
                column=col_index).value = row[col]
        index += 1

def write_to_excel(dataframe, file_name, starting_cell='A1', workbook=None, sheet_name='data'):
    # Create a new workbook
    if workbook is None:
        workbook = px.Workbook()

    wb = workbook

    # Calculate the number of empty rows and columns
    start_col = ord(starting_cell[0].upper()) - ord('A')
    start_row = int(starting_cell[1:]) - 1

    # Convert the DataFrame to a list of lists and add the column headers
    list_df = dataframe.rows()
    list_df = [list(row) for row in list_df]
    data = [dataframe.columns] + list_df

    # Add empty strings for the cells before the starting cell
    
    data = [[''] * start_col + row for row in data]
    data = [[''] * len(data[0])] * start_row + data

    # Add the data to the workbook
    wb.new_sheet(sheet_name, data=data)
    

    # Save the workbook to a file
    wb.save(file_name)


def get_workbook():
    return px.Workbook()

